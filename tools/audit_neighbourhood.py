#!/usr/bin/env python3
"""Check the prepared neighbourhood tables against the raw downloads, value by value.

Usage:   python3 tools/audit_neighbourhood.py
         python3 tools/audit_neighbourhood.py --raw raw-indicators --out work/neighbourhood
         python3 tools/audit_neighbourhood.py --show E00000001 E01000001 S01006646   # what the downloads say for these areas

Run it after tools/prep_neighbourhood.py, and again whenever a download or a table changes. It needs
only openpyxl (pip install openpyxl). Exit code 0 = nothing wrong, 1 = at least one problem.

It is a second, independent look, not a re-run of the first: it reads the downloads with its own
code (plain csv and openpyxl, no pandas) and shares none of prep_neighbourhood.py's logic, so a mistake
in one is unlikely to be repeated in the other. It checks that

  * every area is present, and every value that the download publishes (OAC and LOAC groups, AHAH and
    deprivation ranks, the AHAH score, the AHAH percentile) is exactly what the table holds;
  * the values prep_neighbourhood.py works out itself (deciles, and the deprivation percentile) follow
    the rule ceil(10 or 100 * rank / areas in the country), and where the download publishes a decile
    (England, Wales), how far they agree with it;
  * the direction of each scale is right, by printing the most and least deprived places (and the
    healthiest and least healthy) for you to judge;
  * the files hold plain numbers, no empty cells and no repeated area codes, and are still the files
    that manifest.json recorded.

What is NOT checked here: anything in the TRE. That is what `python3 -m pipeline.nbhd_tables check` and
`lookup` are for; --show is the other half of a spot check: take the area codes that `lookup` printed for a
postcode in the TRE and see here what the download itself says for them.
"""
import argparse
import csv
import glob
import hashlib
import json
import re
import sys
import warnings
from collections import Counter
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
INT, NUMBER = re.compile(r"^-?\d+$"), re.compile(r"^-?\d+(\.\d+)?$")


def ceil_div(a, b):
    return -(-a // b)


class Audit:
    def __init__(self, raw, out):
        self.raw, self.out, self.problems = Path(raw), Path(out), []

    def check(self, ok, message):
        if not ok:
            self.problems.append(message)
        return ok

    def same_areas(self, label, ours, raw):
        """Record a problem, naming some of them, if the table and the download do not hold the same areas.
        Returns the areas both have, so the value checks that follow cannot fail on a missing one."""
        missing, extra = sorted(set(raw) - set(ours)), sorted(set(ours) - set(raw))
        self.check(not missing and not extra, f"{label}: {len(missing)} areas of the download are not in the table "
                   f"(e.g. {missing[:3]}), and {len(extra)} areas in the table are not in the download (e.g. {extra[:3]})")
        return sorted(set(raw) & set(ours))

    def prepared(self, name):
        with open(self.out / f"nbhd_{name}.csv", newline="", encoding="utf-8") as f:
            return list(csv.DictReader(f))

    # -- the four products ------------------------------------------------------------------

    def oac(self):
        print("=== UK OAC: is every value identical to the download?")
        raw = self._oac_raw()
        gb = {k: v for k, v in raw.items() if k[0] in "EWS"}
        ours = {r["area_code"]: (r["oac_supergroup"], r["oac_group"], r["oac_subgroup"]) for r in self.prepared("oac")}
        common = self.same_areas("OAC", ours, gb)
        different = sum(ours[k] != gb[k] for k in common)
        self.check(not different, f"OAC: {different} areas have a value that differs from the download")
        print(f"download {len(raw):,} areas (Northern Ireland {len(raw) - len(gb):,}, left out on purpose); table {len(ours):,}; "
              f"areas whose supergroup, group or subgroup differs: {different}")

    def loac(self):
        print("\n=== London OAC")
        raw = self._loac_raw()
        ours = {r["area_code"]: (r["loac_supergroup"], r["loac_group"]) for r in self.prepared("loac")}
        common = self.same_areas("LOAC", ours, raw)
        different = sum(ours[k] != raw[k] for k in common)
        self.check(not different, f"LOAC: {different} areas have a value that differs from the download")
        print(f"download {len(raw):,}, table {len(ours):,}; areas that differ: {different}")

    def ahah(self):
        print("\n=== AHAH v5.1")
        raw = self._ahah_raw()
        ours = {r["area_code"]: r for r in self.prepared("ahah")}
        n = len(raw)
        common = self.same_areas("AHAH", ours, raw)
        ranks = all(int(ours[k]["ahah_rank"]) == raw[k][1] for k in common)
        pcts = all(int(ours[k]["ahah_pctile"]) == raw[k][2] for k in common)
        self.check(ranks and pcts, "AHAH: a rank or a published percentile differs")
        worst = max(abs(float(ours[k]["ahah_score"]) - raw[k][0]) for k in common)
        self.check(worst <= 0.0005 + 1e-9, "AHAH: a score differs by more than rounding to 3 decimals")
        print(f"{n:,} areas. Rank identical to the download: {ranks}. Published percentile identical: {pcts}. "
              f"Largest difference in the score (rounded to 3 decimals here): {worst:.4f}")
        mine = {k: ceil_div(raw[k][1] * 10, n) for k in raw}
        self.check(all(int(ours[k]["ahah_decile"]) == mine[k] for k in common), "AHAH: a decile is not ceil(10 * rank / areas)")
        own = sum(raw[k][2] == ceil_div(raw[k][1] * 100, n) for k in raw)
        via = sum(ceil_div(raw[k][2], 10) == mine[k] for k in raw)
        print(f"AHAH publishes no decile; ours is ceil(10 * rank / {n:,}). The file's own percentile follows ceil(100 * rank / {n:,}) "
              f"for {own / n:.2%} of areas, and a decile taken from that percentile agrees with ours for {via / n:.2%}.")
        ordered = [mine[k] for k in sorted(raw, key=lambda k: raw[k][1])]
        self.check(all(a <= b for a, b in zip(ordered, ordered[1:])), "AHAH: deciles go backwards as the rank rises")
        print("areas per decile:", dict(sorted(Counter(ordered).items())))

    def deprivation(self):
        print("\n=== Deprivation, country by country")
        ours_all = {r["area_code"]: r for r in self.prepared("imd")}
        for label, reader in (("England IoD2025", self._england), ("Wales WIMD2025", self._wales), ("Scotland SIMD2020v2", self._scotland)):
            raw = reader()
            n, letter = len(raw), next(iter(raw))[0]
            ours = {k: v for k, v in ours_all.items() if k[0] == letter}
            common = self.same_areas(label, ours, raw)
            self.check(sorted(v["rank"] for v in raw.values()) == list(range(1, n + 1)), f"{label}: the download's ranks are not 1..{n}")
            ranks = all(int(ours[k]["imd_rank"]) == raw[k]["rank"] for k in common)
            rule = all(int(ours[k]["imd_decile"]) == ceil_div(raw[k]["rank"] * 10, n)
                       and int(ours[k]["imd_pctile"]) == ceil_div(raw[k]["rank"] * 100, n)
                       and int(ours[k]["imd_areas"]) == n for k in common)
            self.check(ranks, f"{label}: a rank differs from the download")
            self.check(rule, f"{label}: a decile or percentile does not follow ceil(10 or 100 * rank / {n})")
            print(f"\n{label}: {n:,} areas. Ranks identical to the download: {ranks}. Deciles and percentiles follow the rule: {rule}")
            sample = next(iter(raw.values()))
            if "decile" in sample:
                differ = [(k, raw[k]["rank"], raw[k]["decile"], int(ours[k]["imd_decile"]))
                          for k in common if raw[k]["decile"] != int(ours[k]["imd_decile"])]
                print(f"   published decile against ours: {n - len(differ):,} of {n:,} identical, {len(differ)} differ"
                      + ("".join(f"\n     {k} rank {r}: published {p}, ours {m}" for k, r, p, m in differ)))
            else:
                print("   the download has no decile, quintile or percentile, so there is nothing to compare ours with (only the ranks)")
            by_rank = sorted(raw, key=lambda k: raw[k]["rank"])
            print("   rank 1, should be very deprived:", [raw[k]["name"] for k in by_rank[:3]])
            print(f"   rank {n:,}, should be very affluent:", [raw[k]["name"] for k in by_rank[-2:]])

    def direction_of_ahah(self):
        print("\n=== AHAH direction, against place names from the England deprivation file (same LSOA codes)")
        names = self._england()
        ranked = sorted((r for r in self.prepared("ahah") if r["area_code"] in names), key=lambda r: int(r["ahah_rank"]))
        print("healthiest (rank 1 on):", [names[r["area_code"]]["name"] for r in ranked[:3]])
        print("least healthy (last ranks):", [names[r["area_code"]]["name"] for r in ranked[-3:]])

    def files(self):
        print("\n=== The files as the TRE will read them")
        numbers = {"oac": {}, "loac": {}, "ahah": {"ahah_score": NUMBER, "ahah_rank": INT, "ahah_pctile": INT, "ahah_decile": INT},
                   "imd": {"imd_rank": INT, "imd_areas": INT, "imd_pctile": INT, "imd_decile": INT}}
        for name, spec in numbers.items():
            rows = self.prepared(name)
            empty = sum(v == "" for r in rows for v in r.values())
            malformed = sum(not rx.match(r[c]) for r in rows for c, rx in spec.items())
            repeated = len(rows) - len({r["area_code"] for r in rows})
            self.check(not (empty or malformed or repeated), f"nbhd_{name}: empty cells {empty}, malformed numbers {malformed}, repeated codes {repeated}")
            for column in (c for c in spec if c.endswith("decile")):
                self.check({int(r[column]) for r in rows} == set(range(1, 11)), f"nbhd_{name}: {column} is not 1 to 10")
            print(f"nbhd_{name}: {len(rows):,} rows; empty cells {empty}; numbers not in plain form {malformed}; repeated area codes {repeated}")
        manifest = json.load(open(self.out / "manifest.json"))
        same = all(hashlib.sha256((self.out / e["file"]).read_bytes()).hexdigest() == e["sha256"] for e in manifest.values())
        self.check(same, "a file no longer matches the checksum in manifest.json")
        print("files match the checksums in manifest.json:", same)

    def show(self, codes):
        """What the downloads say for some area codes (copied from `python3 -m pipeline.nbhd_tables lookup`
        in the TRE), to compare with what the tables there gave for the same postcode."""
        oac, loac, ahah = self._oac_raw(), self._loac_raw(), self._ahah_raw()
        deprivation = {"England IoD2025": self._england(), "Wales WIMD2025": self._wales(), "Scotland SIMD2020v2": self._scotland()}
        print("What the downloads say (not the prepared tables), for each area code:")
        for code in (c.strip().upper() for c in codes):
            print(f"\n{code}")
            found = False
            if code in oac:
                print("  UK OAC 2021/22   supergroup %s, group %s, subgroup %s" % oac[code])
                found = True
            if code in loac:
                print(f"  London OAC       group {loac[code][1]}")
            elif re.fullmatch(r"E00\d{6}", code):
                print("  London OAC       not in the download (it covers London only)")
            if code in ahah:
                score, rank, pct = ahah[code]
                print(f"  AHAH v5.1        rank {rank:,}, published percentile {pct}, score {score:.3f}; decile by our rule "
                      f"ceil(10 * rank / {len(ahah):,}) = {ceil_div(rank * 10, len(ahah))}   (1 = healthiest)")
                found = True
            for label, data in deprivation.items():
                if code in data:
                    area, n = data[code], len(data)
                    print(f"  {label}: {area['name']}: rank {area['rank']:,} of {n:,}; by our rule decile {ceil_div(area['rank'] * 10, n)}, "
                          f"percentile {ceil_div(area['rank'] * 100, n)}; "
                          + (f"published decile {area['decile']}" if "decile" in area else "no published decile (ranks only)") + "   (1 = most deprived)")
                    found = True
            if not found:
                print("  in none of the downloads. Check the code: Scottish 2022 data zones are only in AHAH, 2011 ones only in "
                      "the deprivation download, and Northern Ireland is not covered.")

    # -- reading the downloads, independently of prep_neighbourhood.py -----------------------

    def _oac_raw(self):
        return {r["GeographyCode"]: (r["Supergroup"], r["Group"], r["Subgroup"])
                for r in csv.DictReader(open(self.raw / "oac21/uk_oac_final.csv", encoding="utf-8-sig"))}

    def _loac_raw(self):
        return {r["OA"]: (r["SG"], r["G"]) for r in csv.DictReader(open(self.raw / "loac21/loac_groups.csv", encoding="utf-8-sig"))}

    def _ahah_raw(self):
        return {r["lsoa21cd"]: (float(r["ahah"]), int(r["ahah_rnk"]), int(r["ahah_pct"]))
                for r in csv.DictReader(open(self.raw / "ahah/ahah_v5_1.csv", encoding="utf-8-sig"))}

    def _england(self):
        sheet = openpyxl.load_workbook(self.raw / "imd/File_1_IoD2025_Index_of_Multiple_Deprivation.xlsx", read_only=True, data_only=True)["IMD25"]
        rows = sheet.iter_rows(values_only=True)
        next(rows)
        return {r[0]: {"rank": int(r[4]), "decile": int(r[5]), "name": f"{r[1]} ({r[3]})"} for r in rows if r[0]}

    def _wales(self):
        found = {}
        for r in csv.DictReader(open(glob.glob(str(self.raw / "imd/welsh*.csv"))[0], encoding="utf-8-sig")):
            if r["Domain"] == "WIMD" and re.fullmatch(r"W01\d{6}", r["Area code"]):
                area = found.setdefault(r["Area code"], {"name": r["Area name"]})
                area[r["Data description"]] = float(r["Data values"].replace(",", "").strip())
        return {k: {"rank": int(v["Rank"]), "decile": int(v["Decile"]), "name": v["name"]} for k, v in found.items()}

    def _scotland(self):
        sheet = openpyxl.load_workbook(self.raw / "imd/SIMD+2020v2+-+ranks.xlsx", read_only=True, data_only=True)["SIMD 2020v2 ranks"]
        rows = sheet.iter_rows(values_only=True)
        next(rows)
        return {r[0]: {"rank": int(r[5]), "name": f"{r[1]} ({r[2]})"} for r in rows if r[0]}


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--raw", default=ROOT / "raw-indicators", help="folder with the downloads")
    parser.add_argument("--out", default=ROOT / "work" / "neighbourhood", help="folder with the prepared tables")
    parser.add_argument("--show", nargs="+", metavar="CODE", help="only print what the downloads say for these area codes")
    args = parser.parse_args()
    warnings.filterwarnings("ignore")
    audit = Audit(args.raw, args.out)
    if args.show:
        audit.show(args.show)
        return 0
    audit.oac()
    audit.loac()
    audit.ahah()
    audit.deprivation()
    audit.direction_of_ahah()
    audit.files()
    if audit.problems:
        print("\nPROBLEMS:\n  " + "\n  ".join(audit.problems))
        return 1
    print("\nno problems found")
    return 0


if __name__ == "__main__":
    sys.exit(main())
